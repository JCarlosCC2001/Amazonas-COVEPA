import openpyxl
import json
import time
import urllib.request
import urllib.error

SUPABASE_URL = "https://byywikicdvmmmndhkooe.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImJ5eXdpa2ljZHZtbW1uZGhrb29lIiwicm9sZSI6ImFub24iLCJpYXQiOjE3OTI1MTIxODgsImV4cCI6MjEwODA4ODE4OH0.KnBrlx7KmH6OelDv9k7EstHhuFCTGcgtgarLvHh4epk"

def clean_val(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s == "" or s.lower() == "none" else s

def migrar():
    print("⏳ Leyendo archivo observaciones.xlsx...")
    wb = openpyxl.load_workbook('observaciones.xlsx', data_only=True)
    sheet = wb.active

    records = []
    for r in range(2, sheet.max_row + 1):
        id_obs = clean_val(sheet.cell(r, 1).value)
        if not id_obs:
            continue
        
        row_dict = {
            "id_obs": id_obs,
            "hito": clean_val(sheet.cell(r, 2).value),
            "id_sitio": clean_val(sheet.cell(r, 3).value),
            "protocolo": clean_val(sheet.cell(r, 4).value),
            "subsistema": clean_val(sheet.cell(r, 5).value),
            "detalle": clean_val(sheet.cell(r, 6).value),
            "item": clean_val(sheet.cell(r, 7).value),
            "sub_item": clean_val(sheet.cell(r, 8).value),
            "comentario_inicial": clean_val(sheet.cell(r, 9).value),
            "comentario_complementario": clean_val(sheet.cell(r, 10).value),
            "comentario_final": clean_val(sheet.cell(r, 11).value),
            "acuerdo_subsanacion": clean_val(sheet.cell(r, 12).value),
            "estado_gilat": clean_val(sheet.cell(r, 13).value),
            "estado_covepa": clean_val(sheet.cell(r, 14).value) or "Pendiente",
            "estado_pronatel": clean_val(sheet.cell(r, 15).value),
            "comentario_pronatel": clean_val(sheet.cell(r, 16).value),
            "comentario_covepa": clean_val(sheet.cell(r, 17).value),
            "fecha_subsanacion": None
        }
        records.append(row_dict)

    total = len(records)
    print(f"✅ Total registros cargados del Excel: {total}")

    batch_size = 100
    endpoint = f"{SUPABASE_URL}/rest/v1/observaciones"

    headers = {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates"
    }

    print("🚀 Iniciando carga hacia Supabase...")
    for i in range(0, total, batch_size):
        chunk = records[i:i + batch_size]
        data_bytes = json.dumps(chunk).encode('utf-8')
        req = urllib.request.Request(endpoint, data=data_bytes, headers=headers, method='POST')

        try:
            with urllib.request.urlopen(req) as resp:
                if resp.status in (200, 201):
                    actual = min(i + batch_size, total)
                    porcentaje = round((actual / total) * 100, 1)
                    print(f"[{actual}/{total}] ({porcentaje}%) subidos con éxito...")
        except urllib.error.HTTPError as e:
            err_body = e.read().decode('utf-8')
            print(f"❌ Error en lote {i}-{i+batch_size}: Código HTTP {e.code}")
            print(f"Detalle del error: {err_body}")
            print("\n💡 Tip: Asegúrate de haber ejecutado primero el script SQL para crear la tabla 'observaciones' en Supabase.")
            break
        except Exception as e:
            print(f"❌ Error inesperado: {str(e)}")
            break

        time.sleep(0.05)

    print("\n🏁 Proceso de migración finalizado.")

if __name__ == "__main__":
    migrar()
